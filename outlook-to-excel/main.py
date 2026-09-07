import os
from pathlib import Path
from io import StringIO
from email import policy
from email.parser import BytesParser
from datetime import datetime, timedelta
from copy import copy

import pandas as pd
from openpyxl import load_workbook


EMAIL_FOLDER = Path("emails")

WORKBOOK_PATH = Path("sample_production_workbook.xlsx")

SHEET_NAME = "REPORT"

START_ROW = 211


# CALCULATE BASIC SEDIMENTS & WATER

def calculate_bsw(gross, net):
    return 0 if gross == 0 else round(
        100 - ((net / gross) * 100)
    )


# GET LATEST EMAIL

eml_files = list(EMAIL_FOLDER.glob("*.eml"))

if not eml_files:
    raise Exception("No .eml files found.")

latest_email = max(eml_files, key=os.path.getmtime)

print(f"\nProcessing: {latest_email.name}")


# READ EMAIL HTML

with open(latest_email, "rb") as f:
    msg = BytesParser(policy=policy.default).parse(f)

html = next(
    (
        part.get_content()
        for part in msg.walk()
        if part.get_content_type() == "text/html"
    ),
    None,
)

if html is None:
    raise Exception("No HTML body found.")


# READ REPORT TABLE

df = pd.read_html(StringIO(html))[0]


# REPORT DATE

date_text = str(df.iloc[31, 0])

for fmt in ("%d-%b-%y", "%d/%b/%y"):
    try:
        report_date = datetime.strptime(
            date_text,
            fmt
        ) - timedelta(days=1)
        break
    except ValueError:
        pass
else:
    raise ValueError(f"Unrecognized date format: {date_text}")


# FIELD HELPERS

def get_field(row_index):
    gross = pd.to_numeric(df.iloc[row_index, 1])
    net = pd.to_numeric(df.iloc[row_index, 3])

    return gross, calculate_bsw(gross, net), net


def get_combined_field(row1, row2):
    gross = (
        pd.to_numeric(df.iloc[row1, 1])
        + pd.to_numeric(df.iloc[row2, 1])
    )

    net = (
        pd.to_numeric(df.iloc[row1, 3])
        + pd.to_numeric(df.iloc[row2, 3])
    )

    return gross, calculate_bsw(gross, net), net


# OPEN WORKBOOK

try:
    wb = load_workbook(WORKBOOK_PATH)
except PermissionError:
    raise Exception("Workbook is open")

ws = wb[SHEET_NAME]


# FIND NEXT EMPTY ROW

next_row = START_ROW

while ws.cell(next_row, 8).value not in (None, ""):
    next_row += 1

print(f"\nWriting to row {next_row}")


# COPY FORMATTING OF PREVIOUS ENTRIES IN THE ALREADY EXISTING WORKBOOK

source_row = next_row - 1

for col in range(1, ws.max_column + 1):

    src = ws.cell(source_row, col)
    dst = ws.cell(next_row, col)

    if src.has_style:
        dst._style = copy(src._style)

    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = copy(src.number_format)


# COLUMN MAPPINGS

date_cols = [8, 13, 18, 23, 28, 33, 38, 43, 48, 53]

gross_cols = [9, 14, 19, 24, 29, 34, 39, 44, 49, 54]


# WRITE DATES

for col in date_cols:
    ws.cell(next_row, col, report_date)

ws.cell(next_row, 58, report_date)
ws.cell(next_row, 63, report_date)


# PRODUCTION FIELDS

fields = [
    get_field(3),      # FIELD_A
    get_field(4),      # FIELD_B
    get_field(5),      # FIELD_C
    get_field(6),      # FIELD_D
    get_field(9),      # FIELD_E
    get_field(10),     # FIELD_F
    get_combined_field(9, 10),  # FIELD_G
    get_field(13),     # FIELD_H
    get_field(14),     # FIELD_I
    get_field(15),     # FIELD_J
]

for (gross, bsw, net), gross_col in zip(fields, gross_cols):

    ws.cell(next_row, gross_col, gross)
    ws.cell(next_row, gross_col + 1, bsw)
    ws.cell(next_row, gross_col + 2, net)


# DELIVERY SECTION A

section_a_gross = pd.to_numeric(df.iloc[24, 1])
section_a_net = pd.to_numeric(df.iloc[24, 3])

section_a_bsw = calculate_bsw(section_a_gross, section_a_net)

ws.cell(next_row, 59, section_a_gross)
ws.cell(next_row, 60, section_a_bsw)
ws.cell(next_row, 61, section_a_net)


# DELIVERY SECTION B

section_b_gross = pd.to_numeric(df.iloc[25, 1])
section_b_net = pd.to_numeric(df.iloc[25, 3])

section_b_bsw = calculate_bsw(section_b_gross, section_b_net)

ws.cell(next_row, 64, section_b_gross)
ws.cell(next_row, 65, section_b_bsw)
ws.cell(next_row, 66, section_b_net)

# SAVE

wb.save(WORKBOOK_PATH)

print(f"\nReport Date : {report_date:%d-%b-%y}")
print(f"Updated Row : {next_row}\n")